"""
pptx_deck.py -- presentation of the essential changes (training draft).

Layout: title, overview/statistics, parameter-value changes, then per top chapter
the key changes; table/figure value changes and references appended compactly.
"""
from __future__ import annotations
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

DARK = RGBColor(0x1F, 0x3A, 0x5F)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GREEN = RGBColor(0x15, 0x80, 0x3D)
ORANGE = RGBColor(0x9A, 0x34, 0x12)
GREY = RGBColor(0x6B, 0x72, 0x80)


def _title_slide(prs, title, sub):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(9), Inches(1.6))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(36)
    r.font.bold = True
    r.font.color.rgb = DARK
    p2 = tb.text_frame.add_paragraph()
    r2 = p2.add_run()
    r2.text = sub
    r2.font.size = Pt(18)
    r2.font.color.rgb = GREY
    return s


def _bullet_slide(prs, title, bullets, title_color=DARK):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9.2), Inches(0.9))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = title_color
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.25), Inches(9.0), Inches(5.6))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for text, lvl, color, bold in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        r = p.add_run()
        r.text = text
        r.font.size = Pt(20 - 3 * lvl)
        r.font.bold = bold
        if color:
            r.font.color.rgb = color
    return s


def _chapter_priority(ch, deut) -> float:
    p = 0.0
    rel = (deut or {}).get("training_relevance")
    p += {"high": 30, "medium": 12, "low": 0}.get(rel, 6)
    for c in ch["changes"]:
        if c.get("kennwerte", {}).get("changed"):
            p += 4
        if (c.get("modality") or {}).get("shift") == "verschaerft":
            p += 3
        if c["kind"] == "new" and (c.get("modality") or {}).get("new") in ("muss", "darf_nicht"):
            p += 2
    if ch.get("part_changed"):
        p += 8
    if ch["mode"] == "new":
        p += 6
    if (ch.get("level") or 3) <= 2:
        p += 3
    return p


def build_pptx(synopse: dict, deutung: dict | None, statistics: dict,
               out_path: str | Path, pair_label: str, top_n: int = 22):
    from ..stages.deutung import label as _lbl   # renders the neutral enums in the doc language
    lang = (deutung or {}).get("language", "de")
    deut_by_id = {d.get("section_id"): d for d in (deutung or {}).get("chapters", [])}
    prs = Presentation()

    _title_slide(prs, f"Änderungen {pair_label}",
                 "Automatisch erzeugte Schulungs-Rohfassung — normpare "
                 "(deterministischer Vergleich + KI-Deutung)")

    # overview
    comp = statistics["comparison"]
    ck = comp["change_kinds"]
    _bullet_slide(prs, "Überblick: Was hat sich geändert?", [
        (f"{ck.get('similar', 0)} geänderte Absätze in {comp['n_chapters_with_changes']} Kapiteln", 0, None, True),
        (f"{ck.get('new', 0)} neue Absätze — {ck.get('removed', 0)} entfallene Absätze", 0, None, False),
        (f"{ck.get('moved_in', 0)} verschobene Inhalte (Umstrukturierung)", 0, None, False),
        (f"{len(comp['kennwert_changes'])} Kennwert-/Grenzwert-Änderungen", 0, ORANGE, True),
        (f"Verbindlichkeit: {comp['modality_shifts'].get('verschaerft', 0)} verschärft, "
         f"{comp['modality_shifts'].get('gelockert', 0)} gelockert", 0, RED, False),
        ("Dokumentumfang: "
         f"{statistics['old']['paragraphs']} → {statistics['new']['paragraphs']} Absätze, "
         f"{statistics['old']['tables']} → {statistics['new']['tables']} Tabellen, "
         f"{statistics['old']['formulas']} → {statistics['new']['formulas']} Formeln", 0, GREY, False),
    ])

    # parameter values
    kws = comp["kennwert_changes"][:14]
    if kws:
        _bullet_slide(prs, "Geänderte Kennwerte und Grenzwerte",
                      [(f"{k['chapter']}:  {k['old']}  →  {k['new']}", 0, ORANGE, False) for k in kws]
                      + ([(f"… und {len(comp['kennwert_changes']) - 14} weitere (siehe Synopse)",
                           0, GREY, False)] if len(comp["kennwert_changes"]) > 14 else []))

    # top chapters
    ranked = sorted((ch for ch in synopse["chapters"]
                     if any(c["kind"] != "cosmetic" for c in ch["changes"])),
                    key=lambda ch: -_chapter_priority(
                        ch, deut_by_id.get(ch.get("new_id") or ch.get("old_id"))))
    for ch in ranked[:top_n]:
        cid = ch.get("new_id") or ch.get("old_id")
        d = deut_by_id.get(cid)
        bullets = []
        if d and d.get("change_overview"):
            bullets.append((d["change_overview"], 0, None, False))
        n_shown = 0
        for ci, c in enumerate(ch["changes"]):
            if c["kind"] == "cosmetic" or n_shown >= 5:
                continue
            dd = None
            if d:
                for x in d.get("interpretations", []):
                    if x.get("change_index") == ci:
                        dd = x
                        break
            if dd and dd.get("change"):
                col = RED if dd.get("obligation") == "tightened" else None
                bullets.append(("• " + dd["change"], 1, col, False))
                if dd.get("impact"):
                    bullets.append(("→ " + dd["impact"], 2, GREY, False))
                n_shown += 1
            elif c.get("kennwerte", {}).get("changed"):
                for k in c["kennwerte"]["changed"]:
                    bullets.append((f"• Kennwert: {k['old']['raw']} → {k['new']['raw']}", 1, ORANGE, True))
                n_shown += 1
            elif c["kind"] == "new" and (c.get("modality") or {}).get("new") in ("muss", "darf_nicht"):
                bullets.append(("• NEU (Pflicht): " + (c.get("new_text") or "")[:160], 1, GREEN, False))
                n_shown += 1
        # add table/figure changes (parameter/limit values) compactly
        if d:
            def _clip(t, n=150):
                t = (t or "").strip()
                return t if len(t) <= n else t[:n] + " …"
            for a in (d.get("tables") or [])[:3]:
                st = a.get("status") or ""
                bullets.append(("▦ " + _clip((a.get("table") or "") + " [" + st + "]: "
                                             + (a.get("change") or ""), 170), 1, ORANGE, False))
                for kw in (a.get("value_changes") or [])[:2]:
                    bullets.append(("– " + _clip(kw, 150), 2, ORANGE, False))
            for a in (d.get("figures") or [])[:2]:
                st = a.get("status") or ""
                bullets.append(("▣ " + _clip((a.get("figure") or "") + " [" + st + "]: "
                                             + (a.get("change") or ""), 170), 1, None, False))
        if d and d.get("practical_note"):
            bullets.append(("Praxis: " + d["practical_note"], 0, DARK, True))
        if not bullets:
            continue
        rel = (d or {}).get("training_relevance")
        title = f"{cid} — {(ch.get('title') or '')[:60]}"
        if rel:
            title += f"   [{_lbl(rel, lang)}]"
        _bullet_slide(prs, title, bullets)

    # references
    tops = statistics["new"]["top_external_refs"][:12]
    _bullet_slide(prs, "Parallel anzuwendende Regelwerke (Top-Referenzen neu)",
                  [(f"{r} ({n}×)", 0, None, False) for r, n in tops])

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
