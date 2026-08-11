"""
docx_synopse.py -- classic tabular synopsis as a Word document.

One table block per changed chapter: old text | new text | change. Table/figure
interpretations from the AI stage are appended as notes where available.
"""
from __future__ import annotations
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

RED = RGBColor(0xB9, 0x1C, 0x1C)
GREEN = RGBColor(0x15, 0x80, 0x3D)
BLUE = RGBColor(0x1E, 0x40, 0xAF)
GREY = RGBColor(0x6B, 0x72, 0x80)
ORANGE = RGBColor(0x9A, 0x34, 0x12)

_KIND_LABEL = {"similar": "geändert", "split": "aufgeteilt", "merged": "zusammengeführt",
               "new": "NEU", "removed": "ENTFALLEN",
               "moved_in": "hierher verschoben", "moved_away": "verschoben"}


def _shade(cell, hexcolor: str):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _clip(t, n=1400):
    t = (t or "").strip()
    return t if len(t) <= n else t[:n] + " […]"


def build_docx_synopse(synopse: dict, deutung: dict | None, out_path: str | Path,
                       pair_label: str, include_cosmetic: bool = False):
    from ..stages.deutung import label as _lbl   # renders the neutral enums in the doc language
    lang = (deutung or {}).get("language", "de")
    # keyed by mapping id (ENT-24), with section_id as the fallback for older runs
    deut_by_mapping, deut_by_section = {}, {}
    if deutung:
        for d in deutung.get("chapters", []):
            if d.get("mapping_id"):
                deut_by_mapping[d["mapping_id"]] = d
            deut_by_section[d.get("section_id")] = d

    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.5)
        s.right_margin = Cm(1.5)
    doc.styles["Normal"].font.size = Pt(9)

    h = doc.add_heading(f"Synopse {pair_label}", level=0)
    doc.add_paragraph(
        f"Gegenüberstellung {synopse['old_doc']} (alt) ↔ {synopse['new_doc']} (neu). "
        "Automatisch erzeugt mit normpare; nur Kapitel mit inhaltlichen Änderungen. "
        "Kosmetische Abweichungen (Zeichensetzung, Umbrüche, PDF-Artefakte) sind unterdrückt.")

    legend = doc.add_paragraph()
    for txt, col in (("NEU ", GREEN), ("ENTFALLEN ", RED), ("verschoben ", BLUE),
                     ("Kennwert-Änderung", ORANGE)):
        r = legend.add_run("■ " + txt + "  ")
        r.font.color.rgb = col
        r.font.size = Pt(9)

    for ch in synopse["chapters"]:
        subst = [c for c in ch["changes"]
                 if include_cosmetic or c["kind"] != "cosmetic"]
        if not subst and not ch.get("part_changed"):
            continue
        cid = ch.get("new_id") or ch.get("old_id")
        title = f"{cid} — {ch.get('title') or ''}"
        if ch["mode"] == "new":
            title += "   [NEUES KAPITEL]"
        elif ch["mode"] == "removed":
            title += "   [KAPITEL ENTFALLEN]"
        part = ch.get("part") or ""
        if part.startswith("anhang"):
            title += f"   ({'normativ' if part.endswith('normativ') else 'informativ'})"
        hp = doc.add_heading(title, level=2)
        if ch.get("part_changed"):
            r = doc.add_paragraph().add_run("⚠ Status normativ/informativ hat sich geändert!")
            r.font.color.rgb = RED
            r.bold = True

        d = deut_by_mapping.get(ch.get("mapping_id")) or deut_by_section.get(cid)
        if d:
            if d.get("keywords"):
                p = doc.add_paragraph()
                r = p.add_run("Schlagworte: " + " · ".join(d["keywords"]))
                r.italic = True
                r.font.color.rgb = BLUE
                r.font.size = Pt(8)
            if d.get("summary_new") or d.get("summary_old"):
                p = doc.add_paragraph()
                r = p.add_run("Kapitelinhalt: ")
                r.bold = True
                p.add_run(d.get("summary_new") or d.get("summary_old") or "")
            if d.get("change_overview"):
                p = doc.add_paragraph()
                r = p.add_run("Änderungsüberblick: ")
                r.bold = True
                p.add_run(d["change_overview"])

        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        for i, t in enumerate(("Alte Fassung", "Neue Fassung", "Änderung")):
            hdr[i].paragraphs[0].add_run(t).bold = True
            _shade(hdr[i], "DBEAFE")
        widths = (Cm(7.2), Cm(7.2), Cm(3.6))

        deut_map = {}
        if d:
            for dd in d.get("interpretations", []):
                i = dd.get("change_index")
                if isinstance(i, int):
                    deut_map[i] = dd

        for ci, c in enumerate(ch["changes"]):
            if c["kind"] == "cosmetic" and not include_cosmetic:
                continue
            row = tbl.add_row().cells
            row[0].text = _clip(c.get("old_text"))
            row[1].text = _clip(c.get("new_text"))
            info = row[2].paragraphs[0]
            lab = info.add_run(_KIND_LABEL.get(c["kind"], c["kind"]) + "\n")
            lab.bold = True
            lab.font.color.rgb = {"new": GREEN, "removed": RED,
                                  "moved_in": BLUE, "moved_away": BLUE}.get(c["kind"], GREY)
            if c["kind"] == "new":
                _shade(row[1], "F0FDF4")
            elif c["kind"] == "removed":
                _shade(row[0], "FEF2F2")
            for k in (c.get("kennwerte") or {}).get("changed", []):
                r = info.add_run(f"{k['old']['raw']} → {k['new']['raw']}\n")
                r.font.color.rgb = ORANGE
                r.bold = True
            mod = c.get("modality") or {}
            if mod.get("shift") in ("verschaerft", "gelockert"):
                r = info.add_run(f"Verbindlichkeit {mod['shift']} ({mod.get('old')}→{mod.get('new')})\n")
                r.font.color.rgb = RED if mod["shift"] == "verschaerft" else BLUE
            if c.get("moved_from"):
                info.add_run(f"aus {c['moved_from']}\n").font.color.rgb = BLUE
            if c.get("moved_to"):
                info.add_run(f"nach {c['moved_to']}\n").font.color.rgb = BLUE
            dd = deut_map.get(ci)
            if dd:
                r = info.add_run(f"{_lbl(dd.get('semantic_label',''), lang)}: {dd.get('change','')}")
                r.font.size = Pt(8)
                r.font.color.rgb = GREY
                qv = dd.get("cross_reference_note") or ""
                if qv and qv != "renummeriert":
                    r = info.add_run(f"\n↪ {qv}")
                    r.font.size = Pt(8)
                    r.font.color.rgb = BLUE
            for cell, w in zip(row, widths):
                cell.width = w

        # formula/table diff as a note
        fd = ch.get("formulas_diff") or {}
        if fd.get("added") or fd.get("removed"):
            p = doc.add_paragraph()
            r = p.add_run(f"Formeln: {len(fd.get('added', []))} neu, "
                          f"{len(fd.get('removed', []))} entfallen")
            r.italic = True
            r.font.color.rgb = ORANGE
        td = [t for t in ch.get("tables_diff") or [] if t["kind"] != "matched" or not t.get("identical")]
        if td:
            p = doc.add_paragraph()
            n_new = sum(1 for t in td if t["kind"] == "new")
            n_rem = sum(1 for t in td if t["kind"] == "removed")
            n_chg = sum(1 for t in td if t["kind"] == "matched")
            r = p.add_run(f"Tabellen: {n_chg} geändert, {n_new} neu, {n_rem} entfallen")
            r.italic = True
            r.font.color.rgb = ORANGE

        # subject-matter table/figure interpretation (from the AI stage)
        if d:
            _STAT = {"changed": ORANGE, "new": GREEN, "removed": RED}
            for a in d.get("tables") or []:
                p = doc.add_paragraph()
                r = p.add_run(f"▦ Tabelle — {a.get('table','')} [{a.get('status','')}]: ")
                r.bold = True
                r.font.color.rgb = _STAT.get(a.get("status"), ORANGE)
                r.font.size = Pt(9)
                p.add_run(a.get("change", "")).font.size = Pt(9)
                for kw in a.get("value_changes") or []:
                    rp = doc.add_paragraph()
                    rr = rp.add_run("    • " + kw)
                    rr.font.color.rgb = ORANGE
                    rr.font.size = Pt(8)
                if a.get("impact"):
                    rp = doc.add_paragraph()
                    rr = rp.add_run("    → " + a["impact"])
                    rr.font.size = Pt(8)
                    rr.font.color.rgb = GREY
            for a in d.get("figures") or []:
                p = doc.add_paragraph()
                r = p.add_run(f"▣ Bild — {a.get('figure','')} [{a.get('status','')}]: ")
                r.bold = True
                r.font.color.rgb = _STAT.get(a.get("status"), ORANGE)
                r.font.size = Pt(9)
                p.add_run(a.get("change", "")).font.size = Pt(9)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
