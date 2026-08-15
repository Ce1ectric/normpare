"""Final, human-readable synopsis (Word) built from the AI interpretation.

Unlike the deterministic tabular synopsis, this one reads as a per-chapter narrative: a
bundled change overview, the essential changes grouped into readable entries (from the
LLM's per-change interpretation), a practical hint, keywords and the training relevance.

The data extraction (:func:`_build_entries`) is separated from the Word rendering so it
can be tested without ``python-docx``.
"""
from __future__ import annotations

from pathlib import Path

from .axes import axis_marker

_RELEVANCE_ORDER = {"high": 0, "medium": 1, "low": 2}


def _build_entries(synopse: dict, deutung: dict) -> list[dict]:
    """Turn the synopse + AI interpretation into a list of readable chapter entries.

    One entry per interpreted chapter (chapters without an overview and without change
    interpretations are skipped). Pure function -- no Word dependency.
    """
    titles: dict = {}
    for ch in synopse.get("chapters", []):
        for key in (ch.get("new_id"), ch.get("old_id")):
            if key and key not in titles:
                titles[key] = ch.get("title") or ""

    entries: list[dict] = []
    for ch in (deutung or {}).get("chapters", []):
        overview = (ch.get("change_overview") or "").strip()
        interpretations = ch.get("interpretations") or []
        changes = []
        for d in interpretations:
            text = (d.get("change") or "").strip()
            if not text:
                continue
            changes.append({
                "label": (d.get("semantic_label") or "").strip(),
                "binding": (d.get("obligation") or "").strip(),
                # AP-17: the three model axes behind the label, empty for a run without them
                "marker": axis_marker(d),
                "text": text,
                "impact": (d.get("impact") or "").strip(),
            })
        if not overview and not changes:
            continue
        sid = ch.get("section_id")
        entries.append({
            "section_id": sid,
            "title": titles.get(sid, ""),
            "relevance": (ch.get("training_relevance") or "").strip(),
            "overview": overview,
            "summary_new": (ch.get("summary_new") or "").strip(),
            "changes": changes,
            "practice": (ch.get("practical_note") or "").strip(),
            "keywords": list(ch.get("keywords") or []),
        })
    return entries


def build_final_synopse(synopse: dict, deutung: dict, out_path: str | Path,
                        pair_label: str) -> None:
    """Render the final, LLM-interpreted synopsis as a Word document."""
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    from ..stages.deutung import label as _lbl   # renders the neutral enums in the doc language

    red, orange, blue, grey, green = (
        RGBColor(0xB9, 0x1C, 0x1C), RGBColor(0x9A, 0x34, 0x12), RGBColor(0x1E, 0x40, 0xAF),
        RGBColor(0x6B, 0x72, 0x80), RGBColor(0x15, 0x80, 0x3D))
    label_color = {"new_obligation": orange, "restricted": orange, "tightened": orange,
                   "removed_obligation": blue, "relaxed": blue,
                   "extended": green, "clarified": grey, "equivalent": grey}
    relevance_color = {"high": red, "medium": orange, "low": grey}
    lang = (deutung or {}).get("language", "de")

    entries = _build_entries(synopse, deutung)
    # most training-relevant chapters first, then by section id
    entries.sort(key=lambda e: (_RELEVANCE_ORDER.get(e["relevance"], 3), str(e["section_id"])))

    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(f"Finale Synopse (KI-gedeutet) — {pair_label}", level=0)
    intro = doc.add_paragraph()
    intro.add_run(
        f"Kapitelweise Zusammenfassung der Änderungen {synopse.get('old_doc', 'alt')} "
        f"(alt) → {synopse.get('new_doc', 'neu')} (neu). Modell: "
        f"{(deutung or {}).get('model', '?')}. Nach Schulungsrelevanz geordnet."
    ).italic = True

    def _run(p, text, *, bold=False, italic=False, color=None):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if color is not None:
            r.font.color.rgb = color
        return r

    for e in entries:
        title = f"Kapitel {e['section_id']}" + (f" — {e['title']}" if e["title"] else "")
        doc.add_heading(title, level=1)

        if e["relevance"]:
            p = doc.add_paragraph()
            _run(p, "Schulungsrelevanz: ", bold=True)
            _run(p, _lbl(e["relevance"], lang), bold=True,
                 color=relevance_color.get(e["relevance"], grey))

        if e["overview"]:
            p = doc.add_paragraph()
            _run(p, "Überblick: ", bold=True)
            _run(p, e["overview"])
        elif e["summary_new"]:
            p = doc.add_paragraph()
            _run(p, "Neu: ", bold=True)
            _run(p, e["summary_new"])

        if e["changes"]:
            doc.add_heading("Wesentliche Änderungen", level=2)
            for c in e["changes"]:
                p = doc.add_paragraph(style="List Bullet")
                tag = " · ".join(_lbl(x, lang) for x in (c["label"], c["binding"]) if x)
                if tag:
                    _run(p, f"[{tag}] ", bold=True, color=label_color.get(c["label"], grey))
                # the axes stay in their neutral schema wording: they have no German
                # rendering yet, and a half-translated marker would read as two vocabularies
                if c["marker"]:
                    _run(p, f"{c['marker']} ", color=grey)
                _run(p, c["text"])
                if c["impact"]:
                    sub = doc.add_paragraph()
                    sub.paragraph_format.left_indent = Cm(0.8)
                    _run(sub, "Auswirkung: ", italic=True, color=grey)
                    _run(sub, c["impact"], italic=True, color=grey)

        if e["practice"]:
            p = doc.add_paragraph()
            _run(p, "Praxishinweis: ", bold=True)
            _run(p, e["practice"], italic=True)

        if e["keywords"]:
            p = doc.add_paragraph()
            _run(p, "Schlagworte: " + ", ".join(e["keywords"]), color=grey)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
