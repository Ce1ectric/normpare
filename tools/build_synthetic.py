"""
build_synthetic.py -- turn the synthetic corpus Markdown into the DOCX the pipeline reads.

``tests/fixtures/synthetic/alt.md`` and ``neu.md`` are the source of truth; the DOCX are
derived artifacts and are not tracked. Regenerate them with::

    poetry run python tools/build_synthetic.py

Markdown conventions (see the corpus README):

===============  ==========================================================
``## <no> ...``  section, level 1; the first token is the section number
``### <no> ...`` section, level 2
``- ...``        list item
``# ...``        document title -- becomes a leading paragraph, no section
anything else    paragraph
===============  ==========================================================

Two things the DOCX has to reproduce from a real standard, because the ingest reader
keys on them:

* headings carry the Word styles ``Heading1``/``Heading2``; the section number is *not*
  in the heading text (Word numbers headings automatically), so a renumbered chapter
  keeps its title and the mapping has to work that out from the content.
* list items carry a real ``w:numPr``, otherwise the reader classifies them as ordinary
  paragraphs and the list inheritance case of the corpus would not be exercised.

The numbers in the Markdown are therefore not written into the document -- they are
checked against the numbering the reader will derive from the heading levels, so a
mismatch between the corpus and its own numbering surfaces here and not three stages later.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SYNTHETIC = ROOT / "tests" / "fixtures" / "synthetic"

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_ITEM = re.compile(r"^[-*]\s+(.*)$")


class CorpusError(RuntimeError):
    """Raised when the Markdown violates the conventions of the corpus."""


def parse_markdown(text: str) -> list[tuple[str, ...]]:
    """Markdown -> blocks ``("heading", level, number, title)`` / ``("para"|"item", text)``.

    Paragraphs may be wrapped over several lines; a blank line ends them.
    """
    blocks: list[tuple[str, ...]] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append(("para", " ".join(buffer)))
            buffer.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        head = _HEADING.match(line)
        if head:
            flush()
            level, rest = len(head.group(1)), head.group(2).strip()
            if level == 1:                     # document title, not a section
                blocks.append(("para", rest))
                continue
            number, _, title = rest.partition(" ")
            if not title.strip():
                raise CorpusError(f"heading without a title: {line!r}")
            blocks.append(("heading", level - 1, number, title.strip()))
            continue
        item = _ITEM.match(line)
        if item:
            flush()
            blocks.append(("item", item.group(1).strip()))
            continue
        buffer.append(line.strip())
    flush()
    return blocks


def numbering(blocks) -> list[tuple[str, str]]:
    """``(declared number, derived number)`` per heading.

    The derived number is what ``ingest_docx`` counts from the heading levels; the
    declared one is the token in the Markdown. They must agree.
    """
    counters = [0] * 10
    out = []
    for block in blocks:
        if block[0] != "heading":
            continue
        _, level, declared, _title = block
        counters[level - 1] += 1
        for i in range(level, 10):
            counters[i] = 0
        out.append((declared, ".".join(str(c) for c in counters[:level] if c > 0)))
    return out


def check_numbering(blocks, source: Path) -> None:
    """Raise if a declared section number is not the one the reader will derive."""
    wrong = [f"{declared} (would be {derived})"
             for declared, derived in numbering(blocks) if declared != derived]
    if wrong:
        raise CorpusError(f"{source.name}: section numbers do not follow the heading "
                          f"levels: {', '.join(wrong)}")


def _list_item(paragraph) -> None:
    """Give a paragraph a real numbering property, as a bulleted list in Word has."""
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    num_pr = OxmlElement("w:numPr")
    for tag, value in (("w:ilvl", "0"), ("w:numId", "1")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), value)
        num_pr.append(el)
    paragraph._p.get_or_add_pPr().append(num_pr)


def build_docx(md_path, docx_path) -> Path:
    """Write the DOCX for one Markdown edition and return its path."""
    from docx import Document

    md_path, docx_path = Path(md_path), Path(docx_path)
    blocks = parse_markdown(md_path.read_text(encoding="utf-8"))
    check_numbering(blocks, md_path)

    doc = Document()
    for block in blocks:
        if block[0] == "heading":
            _, level, _number, title = block
            doc.add_paragraph(title, style=f"Heading {level}")
        elif block[0] == "item":
            _list_item(doc.add_paragraph(block[1], style="List Bullet"))
        else:
            doc.add_paragraph(block[1])
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return docx_path


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--src", default=str(SYNTHETIC), help="directory holding alt.md/neu.md")
    ap.add_argument("--dest", default=None, help="destination directory (default: --src)")
    args = ap.parse_args(argv)

    src = Path(args.src)
    dest = Path(args.dest) if args.dest else src
    for name in ("alt", "neu"):
        path = build_docx(src / f"{name}.md", dest / f"{name}.docx")
        print(f"written: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
