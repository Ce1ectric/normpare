"""Front matter and navigation lists must not take a chapter number.

Reading the 2023 edition of the corpus as DOCX put "Begriffe" at chapter 7 instead of 3:
"Anwendungsbeginn", "Vorwort", "Änderungen" and "Einleitung" carry a heading style, so the
reader counted them as chapters 1 to 4 and shifted everything after them. The PDF reader
never had the problem -- it files those titles under ``vorspann``. Since the two disagreed,
the same document produced different chapter numbers depending on its format, and the
mapping between two editions read from different formats compared the wrong chapters.

The documents here are built in the test, so nothing depends on the corpus and no real
standard text is needed.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from normpare.stages.ingest import SKIP_TITLES, VORSPANN_TITLES


def _docx(tmp_path: Path, headings: list[tuple[int, str]], name: str = "d.docx") -> Path:
    """A DOCX with the given ``(level, title)`` headings, one paragraph each."""
    from docx import Document

    doc = Document()
    for level, title in headings:
        doc.add_paragraph(title, style=f"Heading {level}")
        doc.add_paragraph(f"Text zu {title}.")
    ziel = tmp_path / name
    doc.save(str(ziel))
    return ziel


def _sections(tmp_path: Path, headings: list[tuple[int, str]]) -> list[dict]:
    from normpare.stages.ingest.docx import read_docx

    quelle = _docx(tmp_path, headings)
    meta = read_docx(quelle, tmp_path / "out", "d", "Testdokument", "old")
    return meta["sections"]


AUFBAU = [(1, "Vorwort"), (1, "Einleitung"),
          (1, "Anwendungsbereich"), (1, "Normative Verweisungen"),
          (1, "Begriffe und Abkürzungen"), (2, "Begriffe")]


def test_front_matter_does_not_shift_the_chapter_numbers(tmp_path):
    """The core case: with front matter counted, "Begriffe" ends up at 5 instead of 3."""
    nach_id = {s["title"]: s["id"] for s in _sections(tmp_path, AUFBAU)}
    assert nach_id["Anwendungsbereich"] == "1"
    assert nach_id["Normative Verweisungen"] == "2"
    assert nach_id["Begriffe und Abkürzungen"] == "3"
    assert nach_id["Begriffe"] == "3.1"


def test_front_matter_lands_in_vorspann(tmp_path):
    nach_titel = {s["title"]: s for s in _sections(tmp_path, AUFBAU)}
    for titel in ("Vorwort", "Einleitung"):
        s = nach_titel[titel]
        assert s["id"] == f"vorspann.{titel.lower()}"
        assert s["part"] == "vorspann"


def test_front_matter_keeps_its_text(tmp_path):
    """Filing it under vorspann must not drop the content."""
    nach_titel = {s["title"]: s for s in _sections(tmp_path, AUFBAU)}
    absaetze = nach_titel["Vorwort"]["paragraphs"]
    assert absaetze, "the front matter section lost its paragraphs"
    assert "Vorwort" in (absaetze[0].get("n0") or "")


def test_navigation_lists_get_no_section(tmp_path):
    aufbau = [(1, "Inhalt"), (1, "Bilder"), (1, "Tabellen"),
              (1, "Anwendungsbereich")]
    abschnitte = _sections(tmp_path, aufbau)
    titel = {s["title"] for s in abschnitte}
    assert "Inhalt" not in titel and "Bilder" not in titel and "Tabellen" not in titel
    assert {s["title"]: s["id"] for s in abschnitte}["Anwendungsbereich"] == "1"


@pytest.mark.parametrize("titel", sorted(VORSPANN_TITLES))
def test_every_front_matter_title_is_recognised(tmp_path, titel):
    abschnitte = _sections(tmp_path, [(1, titel.capitalize()), (1, "Anwendungsbereich")])
    nach_titel = {s["title"].lower(): s for s in abschnitte}
    assert nach_titel[titel]["part"] == "vorspann"
    assert {s["title"]: s["id"] for s in abschnitte}["Anwendungsbereich"] == "1"


def test_both_readers_share_one_definition():
    """The guard against the two readers drifting apart again.

    The DOCX path disagreed with the PDF path because each carried its own list. The
    lists now live in the package; this test fails if someone re-introduces a local one.
    """
    from normpare.stages.ingest import docx as docx_reader
    from normpare.stages.ingest import pdf as pdf_reader

    assert pdf_reader._VORSPANN is VORSPANN_TITLES
    assert pdf_reader._SKIP_TOC is SKIP_TITLES
    assert docx_reader.VORSPANN_TITLES is VORSPANN_TITLES
    assert docx_reader.SKIP_TITLES is SKIP_TITLES
