"""AP-41 part A/B: every change knows the section it really sits in.

A chapter mapping is named after its head (``11.2``), and until now every view that
showed a single change showed that head. At 4110 four out of ten changes stand in a
subsection of it -- the reader looked up ``11.2`` and the paragraph was in ``11.2.6.7``.

The section is taken from the assignment paragraph -> section that ``norm_doc.json``
makes, never from the shape of the paragraph id: an annex paragraph (``A.p13``), a
synthetic title paragraph (``....p0``) and any future id scheme must not depend on a
string rule.

Everything is built from dicts and mini documents; no run under ``out/`` is read, no LLM
is called, and no value here is a judgement about a standard.
"""
from __future__ import annotations

import copy
import csv

from normpare.report.changes_csv import (
    CSV_COLUMNS,
    DELIMITER,
    ENCODING,
    build_changes_csv,
)
from normpare.report.axes import change_rows
from normpare.report.location import change_location
from normpare.report.stats import pair_stats
from normpare.report.synopse_det import build_docx_synopse
from normpare.stages.deutung import build_chapter_prompt
from normpare.stages.diff import build_synopse, paragraph_sections

#: The seventeen columns of AP-17, in their order. Written out rather than sliced from
#: CSV_COLUMNS, so that a change to the order fails the test instead of travelling with it.
AP17_COLUMNS = (
    "section_id", "mapping_id", "chapter_title", "change_index", "change_kind",
    "structural_operation", "semantic_status", "normative_direction",
    "affected_components", "indeterminate_reason", "semantic_label", "obligation",
    "change", "impact", "evidence", "evidence_ok", "confidence",
)


# --- mini documents ---------------------------------------------------------------

def _para(pid: str, text: str, kind: str = "text") -> dict:
    return {"id": pid, "n0": text, "n1": text, "kind": kind, "modality": "informativ",
            "modality_counts": {}, "refs_internal": [], "refs_external": [], "values": []}


def _section(sid: str, title: str, paras: list[dict]) -> dict:
    return {"id": sid, "title": title, "level": 1, "part": "main",
            "paragraphs": paras, "tables": [], "figures": [], "formulas": []}


def _doc(doc_id: str, sections: list[dict]) -> dict:
    return {"doc_id": doc_id, "title": doc_id, "sections": sections,
            "source": {"format": "docx", "sha256": "0" * 8}}


def _record(old_ids: list[str], new_ids: list[str], links: list[dict]) -> dict:
    return {"old_id": old_ids[0] if old_ids else None,
            "new_id": new_ids[0] if new_ids else None,
            "old_ids": old_ids, "new_ids": new_ids,
            "mapping_id": f"{'+'.join(new_ids)}<{'+'.join(old_ids)}",
            "old_title": "Entkupplungsschutz", "new_title": "Entkupplungsschutz",
            "old_level": 1, "new_level": 1, "old_part": "main", "new_part": "main",
            "match_type": "id+title", "confidence": 1.0,
            "part_changed": False, "para_links": links}


#: The case from Christian's review: the head is 11.2, the paragraph moved from 11.2.5
#: into 11.2.6.7, and the value changed along the way.
OLD_TEXT = ("Die Anschwingzeit des Entkupplungsschutzes darf hoechstens 1 s betragen, "
            "gemessen ab dem Zeitpunkt der Anregung.")
NEW_TEXT = ("Die Anschwingzeit des Entkupplungsschutzes darf hoechstens 0,5 s betragen, "
            "gemessen ab dem Zeitpunkt der Anregung.")


def _corpus(tmp_path):
    """Head 11.2 with a paragraph that sits in 11.2.5 (old) and in 11.2.6.7 (new)."""
    old_doc = _doc("alt", [
        _section("11.2", "Entkupplungsschutz", [_para("11.2.p0", "Es gilt Abschnitt 11.")]),
        _section("11.2.5", "Anschwingzeit", [_para("11.2.5.p0", OLD_TEXT)]),
    ])
    new_doc = _doc("neu", [
        _section("11.2", "Entkupplungsschutz", [_para("11.2.p0", "Es gilt Abschnitt 11.")]),
        _section("11.2.6.7", "Anschwingzeit", [_para("11.2.6.7.p0", NEW_TEXT)]),
    ])
    links = [{"old_ids": ["11.2.p0"], "new_ids": ["11.2.p0"],
              "kind": "identical", "confidence": 1.0},
             {"old_ids": ["11.2.5.p0"], "new_ids": ["11.2.6.7.p0"],
              "kind": "similar", "confidence": 0.95}]
    rec = _record(["11.2", "11.2.5"], ["11.2", "11.2.6.7"], links)
    syn = build_synopse(old_doc, new_doc, [rec], None, tmp_path / "synopse.json",
                        "alt <-> neu")
    return old_doc, new_doc, syn


def _only_change(syn: dict) -> dict:
    changes = [c for ch in syn["chapters"] for c in ch["changes"]]
    assert len(changes) == 1, changes
    return changes[0]


# --- part A: the record carries the section ---------------------------------------

def test_a_change_in_a_subsection_knows_its_section(tmp_path):
    _, _, syn = _corpus(tmp_path)
    ch = syn["chapters"][0]
    assert (ch.get("new_id") or ch.get("old_id")) == "11.2"     # the head is unchanged
    assert _only_change(syn)["section_new"] == "11.2.6.7"


def test_both_sides_carry_their_own_section(tmp_path):
    _, _, syn = _corpus(tmp_path)
    c = _only_change(syn)
    assert (c["section_old"], c["section_new"]) == ("11.2.5", "11.2.6.7")


def test_a_one_sided_change_has_none_on_the_other_side(tmp_path):
    old_doc = _doc("alt", [_section("5", "Anschluss", [_para("5.p0", "Es gilt Abschnitt 5.")])])
    new_doc = _doc("neu", [
        _section("5", "Anschluss", [_para("5.p0", "Es gilt Abschnitt 5.")]),
        _section("5.4.4.1", "Nachweis", [_para("5.4.4.1.p0",
                                               "Der Nachweis ist dem Netzbetreiber "
                                               "vor der Inbetriebnahme vorzulegen.")]),
    ])
    links = [{"old_ids": ["5.p0"], "new_ids": ["5.p0"], "kind": "identical",
              "confidence": 1.0},
             {"old_ids": [], "new_ids": ["5.4.4.1.p0"], "kind": "new", "confidence": 0.0}]
    syn = build_synopse(old_doc, new_doc, [_record(["5"], ["5", "5.4.4.1"], links)], None,
                        tmp_path / "synopse.json", "alt <-> neu")
    c = _only_change(syn)
    assert c["kind"] == "new"
    assert c["section_old"] is None
    assert c["section_new"] == "5.4.4.1"


def test_the_section_comes_from_the_document_not_the_id_string(tmp_path):
    """Paragraph ids that carry no ``.p`` and no section prefix at all."""
    old_doc = _doc("alt", [
        _section("A", "Anhang A", [_para("abs-0001", "Der Anhang A gilt normativ.")]),
        _section("A.3", "Kennwerte", [_para("abs-0002", OLD_TEXT)]),
    ])
    new_doc = _doc("neu", [
        _section("A", "Anhang A", [_para("abs-0001", "Der Anhang A gilt normativ.")]),
        _section("B.4", "Kennwerte", [_para("abs-0099", NEW_TEXT)]),
    ])
    assert paragraph_sections(old_doc)["abs-0002"] == "A.3"
    links = [{"old_ids": ["abs-0001"], "new_ids": ["abs-0001"], "kind": "identical",
              "confidence": 1.0},
             {"old_ids": ["abs-0002"], "new_ids": ["abs-0099"], "kind": "similar",
              "confidence": 0.95}]
    syn = build_synopse(old_doc, new_doc, [_record(["A", "A.3"], ["A", "B.4"], links)],
                        None, tmp_path / "synopse.json", "alt <-> neu")
    c = _only_change(syn)
    assert (c["section_old"], c["section_new"]) == ("A.3", "B.4")


def test_kennwert_changes_carry_the_section(tmp_path):
    old_doc, new_doc, syn = _corpus(tmp_path)
    comparison = pair_stats(syn, old_doc, new_doc)
    entries = comparison["kennwert_changes"]
    assert len(entries) == 1, entries
    entry = entries[0]
    # the new fields
    assert entry["section"] == "11.2.6.7"
    # and the fields AP-41 promised not to move
    assert entry["chapter"] == "11.2"
    assert entry["mapping_id"] == "11.2+11.2.6.7<11.2+11.2.5"
    assert entry["old"].endswith("1 s") and entry["new"].endswith("0,5 s")


# --- part B: the section appears where the reader looks ---------------------------

def test_the_synopsis_shows_the_section_per_change(tmp_path):
    from docx import Document

    _, _, syn = _corpus(tmp_path)
    out = tmp_path / "Synopse.docx"
    build_docx_synopse(syn, None, out, "4110 alt <-> neu")
    doc = Document(str(out))
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert any("11.2.5 → 11.2.6.7" in text for text in cells), cells
    # the block title stays the head of the chapter mapping
    assert any(p.text.startswith("11.2 — ") for p in doc.paragraphs)


def test_the_csv_appends_the_section_columns(tmp_path):
    _, _, syn = _corpus(tmp_path)
    deutung = {"language": "de", "model": "test", "chapters": [
        {"section_id": "11.2", "mapping_id": syn["chapters"][0]["mapping_id"],
         "interpretations": [{"change_index": 0, "semantic_label": "restricted",
                              "obligation": "tightened", "change": "Die Zeit sinkt.",
                              "evidence": "hoechstens 0,5 s", "evidence_ok": True}]},
    ]}
    assert CSV_COLUMNS[:len(AP17_COLUMNS)] == AP17_COLUMNS
    assert CSV_COLUMNS[len(AP17_COLUMNS):] == ("section_old", "section_new")

    out = build_changes_csv(syn, deutung, tmp_path / "Aenderungen.csv")
    with open(out, newline="", encoding=ENCODING) as fh:
        rows = list(csv.reader(fh, delimiter=DELIMITER))
    assert rows[0] == list(CSV_COLUMNS)
    assert rows[1][-2:] == ["11.2.5", "11.2.6.7"]

    row = change_rows(syn, deutung)[0]
    assert change_location(row, row["section_id"]) == "11.2.5 → 11.2.6.7"


def test_the_prompt_is_unchanged(tmp_path):
    """The two new fields are pipeline property: they are never asked about."""
    _, _, syn = _corpus(tmp_path)
    ch = syn["chapters"][0]
    with_fields, sel = build_chapter_prompt(ch, "alt", "neu")
    assert sel == [0]
    stripped = copy.deepcopy(ch)
    for c in stripped["changes"]:
        c.pop("section_old", None)
        c.pop("section_new", None)
        for entry in (c.get("kennwerte") or {}).get("changed", []):
            entry.pop("value_class", None)
    without_fields, _ = build_chapter_prompt(stripped, "alt", "neu")
    assert with_fields == without_fields
